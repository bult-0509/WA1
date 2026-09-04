from pathlib import Path
from copy import deepcopy
import json, re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from research_content import REPORTS, COMMON

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/deliverables'; OUT.mkdir(exist_ok=True)
TEMPLATE=ROOT/'软件开发实践1-项目要求【2026.9】.docx'
SHOTS=ROOT/'docs/screenshots'
F=json.loads((ROOT/'.work/artifacts/findings.json').read_text(encoding='utf-8'))

def runfont(run,size=12,bold=False,cn='宋体'):
    run.font.name='Times New Roman'; run.font.size=Pt(size); run.font.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),cn)
    run.font.highlight_color=None
    run.font.color.rgb=None

def style_doc(d):
    s=d.styles['Normal']; s.font.name='Times New Roman';s.font.size=Pt(12)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    s.paragraph_format.line_spacing=Pt(20);s.paragraph_format.first_line_indent=Pt(24)
    s.paragraph_format.space_after=Pt(0);s.paragraph_format.space_before=Pt(0)
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=sec.bottom_margin=Cm(2.54);sec.left_margin=sec.right_margin=Cm(3.175)

def fresh(cover=False):
    d=Document(TEMPLATE); ps=d.paragraphs
    if cover:
        start=ps[118]._p; end=ps[131]._p
        children=list(d._element.body); a=children.index(start);b=children.index(end)
        for i,node in enumerate(children):
            if (i<a or i>b) and node.tag!=qn('w:sectPr'): d._element.body.remove(node)
        # 使用老师的封面，补足第四名成员而不挤在已有三行中。
        repl={'项目名称：':'项目名称：上海地铁客流时空规律分析与可视化系统',
              '团队成员':'团队成员：学号待填　林睿信（组长）'}
        # 模板成员行的具体空白格式不作为定位依据。
        members=['林睿信（组长）','吴旻昊','朱子墨']
        mi=0
        for p in d.paragraphs:
            if '项目名称' in p.text:p.text='项目名称：上海地铁客流时空规律分析与可视化系统'
            elif '学号' in p.text or '组长' in p.text:
                p.text=('团队成员：' if mi==0 else '　　　　　')+'学号待填　'+members[min(mi,2)];mi+=1
            elif '完成时间' in p.text:
                e=deepcopy(p._p);p._p.addprevious(e)
                from docx.text.paragraph import Paragraph
                p4=Paragraph(e,p._parent);p4.text='　　　　　学号待填　邢雨晨'
                p.text='完成时间：2026年9月4日'
        for p in d.paragraphs:
            for r in p.runs:
                if '学号' in p.text or '项目名称' in p.text or '完成时间' in p.text:runfont(r,12)
                r.font.highlight_color=None
    else:
        for node in list(d._element.body):
            if node.tag!=qn('w:sectPr'):d._element.body.remove(node)
    style_doc(d);return d

def para(d,text,bold=False):
    p=d.add_paragraph();p.paragraph_format.line_spacing=Pt(20);p.paragraph_format.first_line_indent=Pt(24)
    p.paragraph_format.space_after=Pt(0)
    runfont(p.add_run(text),bold=bold);return p

def heading(d,text,level=1):
    p=d.add_paragraph(style='Heading '+str(min(level,3)))
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(20)
    p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True
    runfont(p.add_run(text),14 if level==1 else 12,True,'黑体');return p

def title(d,text):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(16)
    runfont(p.add_run(text),16,True,'黑体')

def caption(d,text):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(6)
    runfont(p.add_run(text),10.5);return p

def table(d,n,label,headers,rows,widths=None):
    cap=caption(d,f'表{n}　{label}');cap.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    total=8306; ws=widths or [total//len(headers)]*len(headers);ws[-1]+=total-sum(ws)
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(total))
    grid=t._tbl.tblGrid
    for child in list(grid):grid.remove(child)
    for w in ws:
        e=OxmlElement('w:gridCol');e.set(qn('w:w'),str(w));grid.append(e)
    for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
    for row in rows:
        for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Pt(ws[j]/20)
            cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(ws[j]))
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(16)
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:runfont(r,10.5,i==0)
    d.add_paragraph();return t

def figure(d,n,label,filename):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(SHOTS/filename),width=Cm(14.65))
    caption(d,f'图{n}　{label}')

def code(d,text):
    for line in text.strip().splitlines():
        p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(14)
        runfont(p.add_run(line),9)

def group_report():
    d=fresh(True);title(d,'上海地铁客流时空规律分析与可视化系统')
    para(d,'选题1：上海地铁数据分析及可视化。项目名称为“沪上流动”。本稿根据2026年9月4日实际数据处理、程序运行和测试结果生成；学号及另一台物理电脑的断网验收尚待补齐。')
    heading(d,'1 团队成员组成及分工')
    para(d,'小组为上午班[01]，组长林睿信，四名成员计划工作量均为25%。表1列出责任安排。按用户要求，程序和文档由OpenAI Codex辅助生成；这里的任务分工是成员复核与答辩责任，不能当作已经完成的人工编码记录。')
    table(d,1,'团队责任分工',['学号','姓名','详细责任','计划占比'],[
        ['待填','林睿信（组长）','课程对照、需求与集成验收、报告整合、提交复核','25%'],
        ['待填','吴旻昊','数据来源、字段字典、清洗质量及复现核验','25%'],
        ['待填','朱子墨','统计分析、聚类评价、图表结论核验','25%'],
        ['待填','邢雨晨','交互体验、异机运行、截图与PPT、AI记录核验','25%']],[1100,1500,4606,1100])
    heading(d,'2 项目开发环境')
    para(d,'开发环境为Windows，IDE/开发工作台为Codex Desktop；Coding Agent为OpenAI Codex。程序使用CPython 3.13.15（Windows x64嵌入式发行版）。文档与幻灯片生成使用Codex提供的独立文档工具环境，它不属于程序运行依赖。表2列出实际安装和运行过的核心库，全部依赖精确版本保存在requirements.txt。')
    table(d,2,'核心开发库',['库','版本','用途'],[['Streamlit','1.49.1','本地Web交互'],['pandas / NumPy','2.3.2 / 2.3.2','表格处理与数值计算'],['SciPy','1.16.1','Spearman相关分析'],['scikit-learn','1.7.1','KMeans与轮廓系数'],['Plotly','6.3.0','离线交互图表及PNG导出'],['PyArrow','21.0.0','Parquet读写'],['Matplotlib','3.10.6','兼容的绘图工具依赖']])
    para(d,'系统只监听127.0.0.1。演示包随附Python、依赖和处理后数据；完整解压后双击启动.vbs，若Windows脚本宿主不可用则使用启动.cmd。开发环境安装依赖后执行python -m streamlit run app.py。运行时不访问在线地图、接口或模型服务。')
    heading(d,'3 数据说明')
    heading(d,'3.1 来源与规模',2)
    para(d,'采用课程指定的MetroFlow数据[1–3]，范围为2017年5月1日至8月31日的上海302座地铁站。下载文件metroData_InOutFlow.csv为3,788,892行、13列、217,442,845字节。每行是某站在10分钟窗口内的聚合客流，不是一个人的原始刷卡记录。课程所述超过7亿条刷卡记录描述来源数据规模，不能写成本项目读入了7亿行明细。原始客流表覆盖123天；正常规律分析排除已知计数缺损的6天，使用117天。')
    para(d,'辅助表包括302行站点资料、123行日历以及2,952行小时天气。数据按CC BY 4.0说明保留来源与署名。站名沿用发布方英文名称，不自行猜译；经纬度和邻接关系用于离线站点地理图，连接线不代表轨道精确走向。天气是城市代表点的观测或再分析资料，不能解释为逐站天气。表3给出字段字典。')
    table(d,3,'原始客流与辅助字段',['字段','类型/单位','解释'],[
        ['date','日期','服务日期，转换为统一日期类型'],['timeslot','整数','10分钟时段编号'],['startTime / endTime','时间','窗口起止，分析覆盖06:00—23:00'],['station','整数ID','关联站点表'],['inFlow / outFlow','非负整数/人次','进站与出站量'],['CinFlow / CoutFlow','非负整数/人次','发布方推断的通勤行程'],['HBOinFlow / HBOoutFlow','非负整数/人次','居家相关其他行程'],['NHBinFlow / NHBoutFlow','非负整数/人次','非居家行程'],['站名、经纬度、邻接站','文本/数值','站点元数据；用于空间展示'],['isWorday','0/1','来源日历字段；0非工作日、1工作日'],['temperature_2m / rain','摄氏度/mm','小时温度与降雨，汇总为日均温度/日总雨量']],[2850,1600,3856])
    heading(d,'3.2 数据质量与保存策略',2)
    para(d,'检查发现：缺失单元格、完全重复、重复键冲突、非法流量、分项和不一致记录均为0；进出均为0的记录有153,301行，保留为合法零值。2017-05-04、05-08、05-09、06-16、06-27、06-28的计数存在来源已说明的缺损，不用普通均值填补，也不作为真实低谷参与正常规律。它们仍保留在质量页面中。')
    para(d,'处理后小时表631,482行，站日表37,146行，保留全部日期及有效标记；另保存全网日表、站点、日历、日天气和聚类结果。页面只读取约8.4MB的处理后数据。原始大文件保留在开发工作区，代码ZIP不包含原始文件及运行环境；README给出重新下载和重算方式。')
    heading(d,'4 需求分析')
    para(d,'课程第1题要求数据预处理、描述性统计、客流特征规律、出行模式与地图可视化。原文“可以进一步”涉及的数据挖掘属于课程可选项，本组在需求阶段选择KMeans站点聚类作为增强。系统面向课程展示与历史探索，不提供实时调度和未来预测。表4将问题、操作和可验收结果对应起来。')
    table(d,4,'分析问题与功能验收',['问题','功能及图表','对应结果证据'],[
        ['A 高峰何时出现？','工作日/非工作日小时曲线、日期热力图','6.2，图2'],['B 哪些站繁忙，在哪里？','离线站点图、日均排名','6.3，图3'],['C 早晚进出如何变化？','进出曲线、方向指数','6.4，图4'],['D 三类行程占多少？','绝对量、加权占比及小时变化','6.4，图5'],['E 天气与客流是否相关？','散点、箱线、分组Spearman','6.5，图6'],['F 哪些站曲线相似？','固定KMeans分组、典型曲线、评价','6.6，图7—8']],[2350,3350,2606])
    para(d,'功能模块结构为：沪上流动 → 项目概览、时间规律、站点空间、出行模式、天气关联、数据质量与说明；出行模式下设进出方向、出行构成、站点聚类。共同能力为日期、日期类型、站点及方向筛选，当前结果CSV导出和图表PNG导出。')
    para(d,'业务流程为：准备公开数据 → 检查与标记 → 小时/日聚合 → 统计和固定聚类 → 本地加载 → 用户筛选 → 图表与解释 → 导出。用户流程为：完整解压演示包 → 双击启动 → 在浏览器看概览 → 选择问题页面 → 调整筛选 → 阅读图表及样本口径 → 下载证据 → 在启动提示窗口点击确定结束服务。')
    para(d,'天气分析固定使用全网进站量，因此禁用站点和方向控件；聚类固定使用完整四个月有效数据，因此禁用日期、日类型和方向控件，选择站点仅用于地图高亮。系统不接受任意上传，不依赖个人电脑绝对路径，不加入OD、账号、爬虫、在线底图或复杂模型。')
    heading(d,'5 系统设计和实现')
    heading(d,'5.1 数据处理与结构',2)
    para(d,'程序按职责拆成data_processing、analysis、clustering、charts四个模块，app.py组织页面，prepare_data.py负责一次性数据准备，launch.py负责本地启动。这样既能把统计函数单独测试，也便于在答辩中从输入追踪到图表，避免把数据处理藏在页面代码里。')
    para(d,'清洗顺序为列结构与类型检查、日期与时段解析、非负整数验证、重复键检查、三类分项加和检查、辅助表唯一性及关联检查。每小时必须具备6个10分钟窗口，每个站日必须具备17小时。正常全网比较采用共同有效日期，防止不同站点分母不同。异常值不机械删除：业务上可能出现的高客流保留；来源已知的缺损日标记排除。没有将缺失补为0，也没有对合法零流量做均值填充。')
    code(d,"""# 先按日期汇总所选站点，再对日期求均值
per_day = hourly.groupby(['date', 'hour'])['inFlow'].sum()
hourly_mean = per_day.groupby('hour').mean()
# 该口径描述所选站点合计的典型小时，而非逐行平均。""")
    heading(d,'5.2 统计方法与图形选择',2)
    para(d,'总量为有效日期的进站或出站人次之和；日均为总量除以有效日期数。统计同时记录中位数、最小值、最大值和峰值日期。小时曲线先按日期汇总，再对同类日期平均。排名采用相同日期集合，避免把累计量误称为日均。折线体现时间变化，热力图同时比较日期与小时，横向条形图便于排名，地理图展示位置，散点和箱线用于关联与分布比较。')
    para(d,'早高峰使用[07:00,09:00)，晚高峰使用[17:00,19:00)。方向指数=(进站−出站)/(进站+出站)，正值表示进站更多，负值表示出站更多；分母为0时不计算。构成比例=某类总人次/三类合计人次，不能直接平均每行比例。天气使用24小时完整日期，日温度取平均、日雨量求和，按工作日和非工作日分别计算Spearman相关系数，避免不同日类型的总体差异掩盖组内关系。')
    heading(d,'5.3 聚类方案与评价',2)
    para(d,'聚类输入为每站的工作日/非工作日、06—22时、进/出站平均曲线，共2×17×2=68项。每种日类型的34项除以该类型全天进出总量，让模型重点比较形状，规模另用排名解释。全部302站具备有效输入。KMeans便于用平均曲线解释分组；固定random_state=42、n_init=10帮助复现，界面不暴露这些开发参数。')
    para(d,'以分3组为初始参照，比较2、3、4、5组，选择轮廓系数最高者；同分时选择较小组数。评价是在同一批四个月数据上的内部评价，不设虚假的分类标签，不使用准确率、F1或回归指标，也不声称存在独立测试集的预测效果。组号没有大小优劣含义。代表站按到本组平均曲线的距离选择，仅用于帮助解释。')
    heading(d,'5.4 交互、异常处理和交付',2)
    para(d,'所有图表在图题中显示当前日期与口径；CSV补充筛选范围、站点、日类型和有效日期数。PNG由浏览器本地生成。地图使用本地站点坐标和邻接资料，不请求在线瓦片。日期未选完整时提示补齐；只有缺损日或筛选后没有有效记录时显示空状态；数据文件缺失、模型结果缺失、端口被占用时给出明确处理提示。')
    para(d,'演示程序采用单进程本地服务和随包Python，不使用Docker、数据库或云部署。启动器只管理它自己创建的子进程，结束时关闭自己的服务。代码ZIP与便携演示包分离，满足源代码ZIP不超过30MB，同时满足用户要求的免安装演示；两者均进入总压缩包。')
    heading(d,'6 结果分析')
    para(d,'以下结果均来自本次真实处理。除明确说明外，范围为2017-05-01至2017-08-31、全部302站、117个有效日期、进站人次。图1—图10为浏览器实际运行截图，截图与源图保存在附件。计数缺损日不参与正常统计，且不会在趋势图中被当作正常低谷连接。')
    heading(d,'6.1 概览与描述性统计',2)
    para(d,'启动后进入项目概览，如图1。有效进站总量为728,813,810人次，日均6,229,177.86人次，中位数6,854,527人次，最小值3,638,016人次，最大值7,691,783人次（2017-05-19）。这些指标描述进站事件，不能解释为去重乘客人数。平均数低于中位数，与非工作日日总量较低的现象一致，但不构成单独的因果解释。')
    figure(d,1,'项目概览的实际运行效果','S01-overview.png')
    heading(d,'6.2 问题A：时间规律',2)
    para(d,'从左侧进入时间规律，保持全网与完整日期范围，如图2。81个有效工作日日均进站6,993,093.59人次，36个非工作日日均4,510,367.47人次。工作日小时峰值在08:00—09:00，平均967,087.85人次；非工作日在17:00—18:00，平均344,204.42人次。说明两类日期的需求节奏不同，不能仅用一条混合平均曲线概括。热力图用于核对高峰是否在多个日期重复出现。')
    figure(d,2,'工作日与非工作日曲线及日期热力图','S02-time.png')
    heading(d,'6.3 问题B：繁忙站点与空间位置',2)
    para(d,'进入站点空间，地图与排名使用同一筛选口径，如图3。日均进站排名前三为People\'s Square（153,843.46人次）、Xujiahui（102,083.57人次）、Jing\'an Temple（101,644.34人次）。地图显示这些高值站的位置，悬停可查看名称和数值，选择站点后可高亮核对。空间集中现象可以提出进一步调查方向，但不能仅凭本站进出量推断OD流向、换乘人数或精确土地用途。')
    figure(d,3,'离线地理分布与繁忙站点排名','S03-stations.png')
    heading(d,'6.4 问题C、D：方向与出行构成',2)
    para(d,'在出行模式中选择进出方向，如图4。用户可对照工作日和非工作日的进出站曲线，并查看07—09时、17—19时的方向指数。全网进出量总体接近，因此站点层面的方向差异更值得观察；不能用全网指数很小推断每个站都没有方向性。选择单站后，曲线与指数会同步更新。')
    figure(d,4,'进出方向曲线与峰段分析入口','S05-direction.png')
    para(d,'切换到出行构成，如图5。有效进站量中，通勤C为233,455,300人次，占32.03%；居家其他HBO为215,063,027人次，占29.51%；非居家NHB为280,295,483人次，占38.46%。三类之和等于进站总量。数量图和比例图分别回答规模与组成，不能只凭比例判断绝对需求大小。这些类别来自数据发布方推断，本组没有重新训练行程分类器。')
    figure(d,5,'三类出行数量与构成比例','S06-composition.png')
    heading(d,'6.5 问题E：天气关联',2)
    para(d,'进入天气关联，如图6，系统固定展示全网进站量与代表点天气。按日类型分开后，工作日日均温度与客流的Spearman系数为−0.490（81天），雨量为−0.226（81天）；非工作日分别为−0.505和−0.451（36天）。在这个历史窗口中，两类天气变量均呈负向关联，但样本数量、季节变化和节假日等因素限制了解释。不能把这些相关系数写成天气导致客流减少的因果效应。')
    figure(d,6,'天气散点和同类日期分布比较','S07-weather.png')
    heading(d,'6.6 增强分析：站点聚类',2)
    para(d,'在出行模式中选择站点聚类，如图7，日期控件变为固定范围。候选2、3、4、5组的轮廓系数分别为0.4269、0.2987、0.2599、0.2053，因此最终选择2组，分别包含175站和127站。图8展示平均曲线与评价视图。第1组工作日早间进站和晚间出站更突出，是可以直接从曲线观察的模式；没有外部土地利用标签时，不直接将其命名为住宅站。')
    figure(d,7,'固定时期站点分组入口','S08-clusters.png')
    figure(d,8,'分组平均曲线及候选组数评价','S09-evaluation.png')
    para(d,'聚类结果说明在当前相对小时曲线中存在可分辨的形态差异。轮廓系数0.427是内部结构指标，既不是42.7%的准确率，也不是对未来的预测能力。分组没有经过独立月份稳定性检验或外部城市功能标签验证，这些属于后续研究方向。')
    heading(d,'6.7 数据质量说明',2)
    para(d,'进入数据质量与说明，如图9，可查看原始行数、字段数、有效日期和六个缺损日。清洗前后的规模及排除原因可由quality.json和prepare_data.py复核。保留异常日的展示，是为了让使用者理解数据边界，避免只展示“干净”的最终曲线而无法解释为什么少了6天。')
    figure(d,9,'数据规模与质量检查','S10-quality.png')
    heading(d,'6.8 导出、测试和异机验收',2)
    para(d,'每个分析表通过“查看数据并导出CSV”展开，图表右上角相机按钮用于PNG导出，操作界面如图10。CSV保留筛选信息，PNG保留图题与统计范围。具体下载文件与验证结果见docs/测试与验收记录.md；该记录会随最终测试更新。')
    figure(d,10,'数据表与CSV导出入口','S11-export.png')
    para(d,'自动化测试覆盖清洗异常、重复冲突、合法零值、天气完整性、均值分母、峰段边界、构成比例、六页和三个模式、站点筛选及空结果。真实浏览器检查补充了地图、日期坐标及页面交互。开发机中文与空格新目录的运行检查属于移目录验证，不能替代另一台物理电脑的断网测试。正式异机验收需在Windows 10/11 x64、无预装Python、普通权限及断网条件下完成六页和导出，并补拍S12。')
    heading(d,'6.9 课程要求与证据对应',2)
    para(d,'表5用于核对课程要求、项目实现、报告章节和最终证据。尚未完成的外部验收保持待验，避免用材料存在代替实际功能通过。')
    table(d,5,'课程要求—项目实现—报告—证据',['课程/用户要求','项目实现','报告位置','验收证据'],[
        ['课程：数据预处理','验证、标记、聚合','3、5.1、6.7','源码、quality.json、图9'],['课程：描述性统计','总量、均值、中位数、极值','5.2、6.1','结果表、图1'],['课程：客流规律','工作日/非工作日比较','6.2','图2、导出CSV'],['课程：地图可视化','离线站点地理图','6.3','图3、断网待补验'],['课程：出行模式','方向与C/HBO/NHB','6.4','图4—5、公式核对'],['用户已选：聚类','KMeans及轮廓系数','5.3、6.6','图7—8、评价CSV'],['用户：天气关联','分组Spearman','6.5','图6、有效天数'],['用户：交互与导出','六页筛选、CSV/PNG','4、6.8','图1—10、测试记录'],['用户：拷贝即运行','内置环境与数据','2、5.4、6.8','便携ZIP；真实异机待验'],['课程：代码交付','README、注释、≤30MB','5.4','代码.zip、包大小记录'],['课程：材料及AI过程','报告、PPT、个人调研、附件','7、8','交付清单、原始记录']],[2000,2150,1500,2656])
    heading(d,'7 AI辅助过程')
    para(d,'实际使用OpenAI Codex。用户首先要求以课程原文为约束进行需求审讯，确认题目1、四人团队、相同工作量和Windows拷贝即用；随后要求剔除不适合本科答辩的复杂参数；最后于2026年9月4日发出“开始按计划书执行”。该授权触发本次开发。需求旧版、精简反馈图、开发前对话副本、首版代码及首次失败日志均有保留。')
    para(d,'首版地图因合并时产生重复站名列而失败，后改为只按ID合并指标；首次真实浏览器检查发现日期热力图的短日期被自动当作年份，后明确使用分类日期轴；分组代表站改为选择接近平均曲线的站点。AI生成代码并未直接被当作正确结果，修复后重新运行测试与界面检查。详细Prompt、前后差异及日志见《AI辅助过程附件》。')
    para(d,'目前有证据的人工贡献是提出需求、确定范围、简化方案和授权实施；程序修改与材料撰写由Codex执行。尚无成员人工改写代码或独立完成调研的记录，不能把计划分工写成已发生事实。四份调研为按成员主题生成的待本人核验稿，成员应阅读引用来源、检查文字并留下实际修改。')
    heading(d,'8 总结')
    para(d,'项目把公开历史客流转为可离线运行的六页系统，完成清洗、时间与空间规律、进出方向、行程构成、天气关联及可解释聚类。主要特色是共同有效日期口径、离线地理图、曲线形态分组和清晰的质量说明；没有声称提出新算法。代码与演示环境分包，便于同时控制提交体积与降低现场运行门槛。')
    para(d,'限制包括数据较早、覆盖时期较短、天气只有一个城市代表点、行程类别来自发布方推断、聚类缺少独立稳定性验证。进一步改进可以围绕独立时期复核、站点外部属性解释和真实使用测试展开。本次提交前还必须补齐学号、完成人员内容核验和另一台电脑的断网验收，并在Word中逐页检查版面。正式发送前核对[01]林睿信_吴旻昊_朱子墨_邢雨晨_地铁命名，截止2026年9月8日23:00；本次未发送邮件。')
    heading(d,'参考资料')
    for i,r in enumerate([COMMON,'Ariza-Sun. MetroFlow[EB/OL]. https://github.com/Ariza-Sun/MetroFlow','MetroFlow dataset[DS/OL]. https://doi.org/10.6084/m9.figshare.28844942','scikit-learn developers. Clustering[EB/OL]. https://scikit-learn.org/stable/modules/clustering.html','课程资料：《软件开发实践1-项目要求【2026.9】》，本地原始DOCX。'],1):para(d,f'[{i}] {r}')
    d.save(OUT/'期末报告.docx')

def individual_reports():
    stats={}
    for item in REPORTS:
        d=fresh();title(d,item['title'])
        para(d,f'《软件开发实践1》个人调研报告　姓名：{item["name"]}　学号：待填　班级：[01]')
        para(d,'版本说明：2026年9月4日由OpenAI Codex按已确定主题辅助生成，供本人阅读来源、核验和修改。未将本稿冒充已经完成的独立人工写作。引用网页查阅日期：2026年9月4日。')
        body=[]
        for h,paragraphs in item['sections']:
            heading(d,h)
            for text in paragraphs:para(d,text);body.append(text)
        heading(d,'参考文献')
        for i,r in enumerate(item['refs'],1):para(d,f'[{i}] {r}')
        count=len(re.findall(r'[\u4e00-\u9fff]',''.join(body)))
        assert count>=1500,(item['name'],count)
        stats[item['name']]={'正文汉字数_不含标题参考文献':count,'参考文献数':len(item['refs'])}
        d.save(OUT/(item['name']+'-调研报告.docx'))
    (ROOT/'.work/artifacts/调研字数检查.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')

def ai_appendix():
    d=fresh();title(d,'AI辅助过程附件')
    para(d,'项目：沪上流动——上海地铁客流时空规律分析与可视化系统。执行工具：OpenAI Codex。记录日期：2026年9月4日。本文是已有证据的整理，不是补造的开发日记。')
    heading(d,'1 真实输入与范围演进')
    for label,text in [('最初任务摘要','用户要求先完整阅读课程DOCX，按选题原文开展需求审讯，形成能够交给Coding Agent的规格，并在开发授权前不进入大规模编码。完整可见消息见对话记录。'),('成员与运行要求原话','“四个成员工作量占比调整为均等”；“windows系统，必须做到拷过去即可运行”。'),('精简要求原话','“在保留高级感的基础上剔除掉部分不适合出现在本科生项目的参数”。'),('开发授权原话','“开始按计划书执行”。')]:
        heading(d,label,2);para(d,text)
    heading(d,'2 AI最初内容与保留方式')
    para(d,'需求v1.0保存在requirements/history中；v1.1按用户意见删减复杂工程结构、多模型评价和秒级指标，保留五个问题、六页、地理图、一个聚类算法和简单评价。开发开始前的对话在本地raw目录备份；交付附件只包含可见用户/助手消息及必要运行记录，不包含内部推理或系统元数据。')
    para(d,'代码首版存为“首版代码-首次界面测试前.zip”。该副本在首轮界面测试之前保存，但数据准备已经执行，因此不将它称为未经任何验证的第一条生成输出。首次失败日志“首次测试-地图列名冲突.log”与最终源码共同保留，以便核对变化。')
    heading(d,'3 关键迭代与修改前后')
    table(d,1,'AI生成、验证、修改与结果',['问题','初始行为','修改','验证依据'],[
        ['地图站名列冲突','合并结果带重复name列，页面失败','仅合并站点ID与数值/组别','原始失败日志、最终测试'],['日期轴识别','05-01等短日期被自动解析成年份','明确分类轴和日期刻度','真实浏览器前后检查'],['分组代表站','先按原表顺序列举站点','按到本组平均曲线距离选代表站','源码、labels.csv'],['参数过多','需求旧版有复杂工程及额外评价','按v1.1精简为KMeans和轮廓系数','需求版本、用户反馈图'],['导出触发重跑','按钮默认触发Streamlit重跑','使用on_click=ignore','浏览器下载复核记录']],[1750,2200,2300,2056])
    para(d,'地图修复体现了分层验证的必要：数据和模型成功运行并不代表页面没有错误。日期轴问题则说明图表能够显示也不等于表达正确。每次修复都对应具体问题，没有为了增加测试数量反复做无关检查。')
    heading(d,'4 人工指导与AI执行的边界')
    para(d,'用户明确选题、成员、工作量、运行方式和项目难度边界，接受默认方案，并授权实施。Codex读取课程、核对数据来源、生成代码、执行分析、测试、修复以及撰写材料。当前没有证据证明成员已经人工编写了某段代码、完成了异机测试或独立改写了调研，因此不填入此类记录。')
    para(d,'后续成员真实参与时，应追加日期、具体文件、修改内容和原因；保留修改前后文本或Git差异。本人阅读参考资料、手工核对样本、检查图表和练习答辩，都是可以如实记录的工作，但必须在实际发生后填写。')
    heading(d,'5 最终成果与初始生成内容的差异')
    para(d,'最终系统从需求候选走向真实数据运行：明确原始表为378万余行聚合记录，正常分析117天，保留6个缺损日解释；得到实际站点排名、类别构成、天气相关和2组聚类；修复了地图关联与日期显示；使用双包交付降低现场安装要求。报告中的结果来自本次计算，不把初始占位文字替换成无依据的漂亮数字。')
    heading(d,'6 证据目录与仍需完成的工作')
    para(d,'证据包括需求历史、用户精简反馈图、可见对话记录、首版代码ZIP、首轮失败日志、最终测试与验收记录、当前源码和实际截图。原始数据与开发环境留在工作区；交付代码包保留轻量结果及重算说明。')
    para(d,'学号、成员对材料的阅读核验、另一台物理Windows电脑的断网运行和最终Word逐页校对仍需完成。不能为满足课程材料要求而伪造这些过程，也不能把本附件当作已经完成的人工贡献证明。')
    d.save(OUT/'AI辅助过程附件.docx')

if __name__=='__main__':
    group_report();individual_reports();ai_appendix()
    print((ROOT/'.work/artifacts/调研字数检查.json').read_text(encoding='utf-8'))
